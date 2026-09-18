import os
import unittest
from app import app, is_web_environment


class WebEnvironmentTests(unittest.TestCase):
    def setUp(self):
        self.client = app.test_client()

    def test_environment_detection_flags(self):
        # Test explicit LOCALPDF_MODE
        orig_mode = os.environ.get("LOCALPDF_MODE")
        orig_railway = os.environ.get("RAILWAY_ENVIRONMENT")
        try:
            os.environ["LOCALPDF_MODE"] = "web"
            self.assertTrue(is_web_environment())

            os.environ["LOCALPDF_MODE"] = "local"
            self.assertFalse(is_web_environment())

            del os.environ["LOCALPDF_MODE"]
            os.environ["RAILWAY_ENVIRONMENT"] = "production"
            self.assertTrue(is_web_environment())
        finally:
            if orig_mode is not None:
                os.environ["LOCALPDF_MODE"] = orig_mode
            elif "LOCALPDF_MODE" in os.environ:
                del os.environ["LOCALPDF_MODE"]

            if orig_railway is not None:
                os.environ["RAILWAY_ENVIRONMENT"] = orig_railway
            elif "RAILWAY_ENVIRONMENT" in os.environ:
                del os.environ["RAILWAY_ENVIRONMENT"]

    def test_security_and_cache_headers_on_convert(self):
        import io
        import fitz

        doc = fitz.open()
        p = doc.new_page()
        p.insert_text((50, 50), "Teste headers")
        pdf_bytes = doc.tobytes()
        doc.close()

        resp = self.client.post(
            "/convert",
            data={"tool": "compress-pdf", "files": (io.BytesIO(pdf_bytes), "doc.pdf")},
            content_type="multipart/form-data",
        )
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.headers.get("X-Content-Type-Options"), "nosniff")
        self.assertEqual(resp.headers.get("X-Frame-Options"), "SAMEORIGIN")
        self.assertIn("no-store", resp.headers.get("Cache-Control", ""))
        self.assertIn("no-cache", resp.headers.get("Pragma", ""))

    def test_about_page_contains_web_vs_desktop_differentiator(self):
        resp = self.client.get("/about")
        self.assertEqual(resp.status_code, 200)
        html = resp.get_data(as_text=True)

        # Verify the comparison section and modalities exist
        self.assertIn("Versão Web (Railway) vs.", html)
        self.assertIn("Aplicativo Desktop (Windows .msi)", html)
        self.assertIn("Versão Web (Nuvem Railway)", html)
        self.assertIn("LocalPDF Web (Railway)", html)
        self.assertIn("Processamento Efêmero", html)
        self.assertIn("Zero Persistência", html)
        self.assertIn("LocalPDF.msi", html)

    def test_web_mode_template_flag(self):
        orig_mode = os.environ.get("LOCALPDF_MODE")
        try:
            os.environ["LOCALPDF_MODE"] = "web"
            resp_web = self.client.get("/")
            html_web = resp_web.get_data(as_text=True)
            self.assertIn("const isServerWebMode = true;", html_web)

            os.environ["LOCALPDF_MODE"] = "local"
            resp_local = self.client.get("/")
            html_local = resp_local.get_data(as_text=True)
            self.assertIn("const isServerWebMode = false;", html_local)
        finally:
            if orig_mode is not None:
                os.environ["LOCALPDF_MODE"] = orig_mode
            elif "LOCALPDF_MODE" in os.environ:
                del os.environ["LOCALPDF_MODE"]


    def test_excel_to_pdf_conversion_quality(self):
        import io
        from datetime import datetime
        import openpyxl
        from openpyxl.worksheet.formula import ArrayFormula
        import pymupdf as fitz

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Janeiro"
        ws["A1"] = "Mês"
        ws["B1"] = "Data"
        ws["C1"] = "Valor"
        ws["A2"] = '="Janeiro "&AnoCalendario'
        ws["B2"] = datetime(2021, 1, 15)
        ws["C2"] = 1500.50
        ws["A3"] = ArrayFormula("A3:C3", "CALC_ROW")

        xlsx_bytes = io.BytesIO()
        wb.save(xlsx_bytes)
        xlsx_bytes.seek(0)

        resp = self.client.post(
            "/convert",
            data={"tool": "excel-to-pdf", "files": [(xlsx_bytes, "calendario.xlsx")]},
            content_type="multipart/form-data",
        )
        self.assertEqual(resp.status_code, 200)
        self.assertIn("calendario.pdf", resp.headers.get("Content-Disposition", ""))

        # Verify generated PDF content
        pdf_doc = fitz.open(stream=resp.data, filetype="pdf")
        self.assertGreaterEqual(pdf_doc.page_count, 1)
        text = pdf_doc[0].get_text()
        pdf_doc.close()

        # Must not contain Python openpyxl object pointers or unformatted timestamps
        self.assertNotIn("<openpyxl.", text)
        self.assertNotIn("00:00:00", text)
        self.assertIn("15/01/2021", text)
        self.assertIn("Janeiro", text)

    def test_robots_txt(self):
        resp = self.client.get("/robots.txt")
        self.assertEqual(resp.status_code, 200)
        self.assertIn("text/plain", resp.headers.get("Content-Type", ""))
        body = resp.get_data(as_text=True)
        self.assertIn("User-agent: *", body)
        self.assertIn("Allow: /", body)
        self.assertIn("Sitemap: https://localpdf.io/sitemap.xml", body)

    def test_sitemap_xml(self):
        resp = self.client.get("/sitemap.xml")
        self.assertEqual(resp.status_code, 200)
        self.assertIn("xml", resp.headers.get("Content-Type", ""))
        body = resp.get_data(as_text=True)
        self.assertIn("<urlset", body)
        self.assertIn("<loc>https://localpdf.io/</loc>", body)
        self.assertIn("<loc>https://localpdf.io/editor</loc>", body)
        self.assertIn("<loc>https://localpdf.io/about</loc>", body)
        self.assertIn("<loc>https://localpdf.io/tool/merge-pdf</loc>", body)
        self.assertIn("<loc>https://localpdf.io/tool/compress-pdf</loc>", body)
        self.assertIn("<loc>https://localpdf.io/tool/pdf-to-word</loc>", body)

    def test_seo_metadata_home(self):
        resp = self.client.get("/")
        self.assertEqual(resp.status_code, 200)
        html = resp.get_data(as_text=True)
        self.assertIn("<title>LocalPDF.io — Ferramentas de PDF 100% Privadas", html)
        self.assertIn('<meta name="description"', html)
        self.assertIn('<meta name="robots" content="index, follow">', html)
        self.assertIn('<link rel="canonical" href="https://localpdf.io/">', html)
        self.assertIn('<meta property="og:title"', html)
        self.assertIn('<meta name="twitter:card" content="summary_large_image">', html)
        self.assertIn('"@type": "WebApplication"', html)
        self.assertIn('"@type": "Organization"', html)

    def test_seo_metadata_tool_pages(self):
        resp = self.client.get("/tool/merge-pdf")
        self.assertEqual(resp.status_code, 200)
        html = resp.get_data(as_text=True)
        self.assertIn("Mesclar PDFs Online Grátis", html)
        self.assertIn('<link rel="canonical" href="https://localpdf.io/tool/merge-pdf">', html)
        self.assertIn('"@type": "BreadcrumbList"', html)
        self.assertIn('"@type": "FAQPage"', html)

        resp_split = self.client.get("/tool/split-pdf")
        self.assertEqual(resp_split.status_code, 200)
        html_split = resp_split.get_data(as_text=True)
        self.assertIn("Dividir PDF Online", html_split)
        self.assertIn('<link rel="canonical" href="https://localpdf.io/tool/split-pdf">', html_split)

    def test_navigation_and_banner_markup(self):
        resp = self.client.get("/")
        self.assertEqual(resp.status_code, 200)
        html = resp.get_data(as_text=True)
        # Check UI back buttons now call navigateBack()
        self.assertIn('onclick="navigateBack()"', html)
        # Check navigateBack and updateDocumentSEO exist in client script
        self.assertIn('function navigateBack()', html)
        self.assertIn('function updateDocumentSEO(', html)
        # Check web mode banner has clean pill elements and dismiss button
        self.assertIn('web-mode-badge-pill', html)
        self.assertIn('dismissWebBanner()', html)
        self.assertIn('web-mode-close-btn', html)


if __name__ == "__main__":
    unittest.main()


